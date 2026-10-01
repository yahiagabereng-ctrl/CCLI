#!/usr/bin/env python3
"""Generate TesPro HW validation Word doc vs CEI 0-16 Annex O/T."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "lab" / "TesPro_HW_Validation_CEI_Annex_O.docx"

ROWS_A3 = [
    ("O.13.1.1.1", "Two IEEE 802.3-2015 external Ethernet IFs: Eth_A (DSO), Eth_B (remote actors).",
     "Which physical ports map to Eth_A vs Eth_B on TR400/TG544/TG-524?", "Port map + panel photo"),
    ("O.13.1.1.1", "Eth_A: 100BaseFX, optical, dual LC, 1310 nm multimode fibre.",
     "Onboard FX port or certified SFP/media converter PN?", "Module datasheet + compliance"),
    ("O.13.1.1.1", "Eth_B: 10Base-T / 100Base-TX, auto-neg, MDI/MDIX; may be SFP (note 2).",
     "Confirm 10/100 vs GE; SFP option for Eth_B?", "PHY/SFP specification"),
    ("O.13.1.1.1", "TCP/IPv4 (mandatory); IPv6 optional; IEC 61850; TLS 1.2+ (62351-3/-8/-9); DHCP; DNS.",
     "Which features enabled on external IFs in product profile?", "Config export / app note"),
    ("O.13.1.1.1", "Physical + data-link status in apparatus log; logical service detectable.",
     "How is link status exposed (driver, SNMP, ubus)?", "API / log sample"),
]

ROWS_A2 = [
    ("O.13.1", "Distinct IFs for DSO, remote actors, plant; plant on separate interfaces.",
     "WAN/LAN1-4/RS485 roles for CCI product?", "Network diagram"),
    ("O.13.1", "No data forwarding between CCI comm interfaces; no switch/bridge on CCI data paths.",
     "Default OpenWrt zones; proof Eth_A not routed to plant LAN?", "Firewall test procedure"),
]

ROWS_A4 = [
    ("O.13.1.1.2", "Local config: serial or USB only — no IP on commissioning port.",
     "Web/SSH on LAN — how locked for CEI product?", "Hardening guide"),
    ("O.13.1.1.2", "User authentication on local access.", "Default auth model?", "Manual section"),
]

ROWS_A5 = [
    ("O.13.1.2", "Eth_A: IEC 61850 per Annex T; signals only via CCI.",
     "Which port is wired to DSO path?", "Install drawing"),
    ("O.13.1.2", "DSO comm loss → autonomous mode (delay per Operating Rule / Annex U).",
     "UPS/battery ≥1 h support (O.13.3)?", "Battery/UPS PN"),
    ("O.13.1.3", "Eth_B: remote enabled actors; Annex T cyber.", "Which port = Eth_B?", "Port map"),
]

ROWS_A6 = [
    ("O.13.1.4", "Plant IF: serial/Ethernet; Modbus RTU/TCP, 61850, etc.; adequate performance.",
     "RS485 isolation rating; max baud; CCLI + seriald sharing?", "RS485 chapter + isolation kV"),
    ("O.13.1.4", "Eth_A and Eth_B must NOT be used for plant internal network.",
     "Confirm plant Modbus only on RS485/plant LAN?", "Routing policy"),
    ("O.13.1.4", "Monitor CCI ↔ plant comm continuously.", "Platform hooks for link monitor?", "—"),
    ("Lab", "TG544: /dev/ttyS2 (A2/B2 screw) I/O error; LuCI RS485-2 → /dev/rs485_2_uart works.",
     "Production mapping: ttyS2 fix or rs485_2_uart only?", "Written confirmation + FW note"),
]

ROWS_A7 = [
    ("O.13.2 / O.13.2.1", "PdC V, P, Q inputs; converter accuracy (class ≤0.2 or ≤5% 100–500 kW PV V5).",
     "Use AI1/AI2 or external meter only?", "Meter / AI spec"),
]

ROWS_A8 = [
    ("O.13.3", "Backup power ≥1 h for CCI and comm equipment.", "Battery option on TG-524?", "UPS datasheet"),
    ("O.13.5", "UTC time; uncertainty ≤ ±100 ms.", "GNSS vs NTP on platform?", "Time sync doc"),
]

ROWS_A9 = [
    ("O.14", "Log DSO/external/plant comm; physical & data-link status; cyber events.",
     "Per-port link logging on TesproOS?", "logread/syslog example"),
]

ROWS_B = [
    ("O.9.2.x", "Active power limitation at PdC (DSO command mandatory PV/wind ≥100 kW — V5).",
     "DO rating for curtail relay on DIO1?", "DIO electrical spec"),
    ("O.13.1", "Plant actuation on separate IF from Eth_A/B.", "DIO as plant-side output — confirm.", "—"),
    ("O.12", "Installation examples; teledistacco predisposition.", "Recommended DIO for Annex M?", "Drawing"),
    ("GD32 / DIO1–4", "Field DI/DO via dido_v2 (lab: DIO1 DO curtail, DIO2 DI permissive).",
     "Voltage, current, isolation; DO state on power-down/reboot?", "GD32 I/O datasheet"),
]

ROWS_C = [
    ("T.3.3.4.10", "Segregation: NAT, VLAN, firewall, VPN; public net via secure VPN only.",
     "On-box segregation vs external router?", "Network architecture PDF"),
    ("T.3.3.4.11", "Secure local commissioning comms.", "Same as O.13.1.1.2.", "—"),
    ("T.3.2.1", "61850 measurements every 4 s; performance Type 3.", "CPU/ETH load guidance?", "—"),
]

ROWS_MIN = [
    ("1", "O.13.1.1.1", "Eth_A 100BaseFX LC 1310 nm", "Module PN?"),
    ("2", "O.13.1.1.1", "Eth_B 10/100TX or SFP", "Which port?"),
    ("3", "O.13.1", "No bridge between IFs", "Firewall proof?"),
    ("4", "O.13.1.4", "Plant RS485 isolation", "ttyS2 vs rs485_2_uart?"),
    ("5", "O.13.1.4", "No plant on Eth_A/B", "Policy doc?"),
    ("6", "O.13.1.1.2", "Local IF no IP", "Product config?"),
    ("7", "O.13", "EMC / isolation", "CE report?"),
    ("8", "O.13.3", "UPS 1 h", "Battery option?"),
    ("9", "DIO / O.9", "Curtail DO spec", "12 V OK?"),
    ("10", "T.3.3.4.10", "Traffic segregation", "On-box or router?"),
]


def add_table(doc: Document, headers: list[str], rows: list[tuple]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = str(val)
            for p in table.rows[ri].cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    title = doc.add_heading("TesPro HW Validation vs CEI 0-16", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        f"Central Plant Controller (CCI) — Annex O / T clause traceability\n"
        f"Supplier: TesPro Electronics · Product: TR400 / TG544 / TG-524\n"
        f"Document ID: CCLI-LAB-TESPRO-HW-VAL-001 · Date: {date.today().isoformat()}\n"
        f"Normative reference: CEI 0-16 consolidated working edition 2025-12 (Allegato O, T)"
    )
    run.font.size = Pt(10)

    doc.add_paragraph(
        "Purpose: Request written hardware validation evidence from TesPro for notified body review. "
        "Software Phase 1 regulation-check (PF2 equations, DIO logic) is separate from this pack."
    )

    doc.add_heading("Part 0 — Supplier response (10 priority items)", level=1)
    add_table(
        doc,
        ["#", "Clause", "Topic", "TesPro response"],
        ROWS_MIN,
    )

    doc.add_heading("Part A — Annex O O.13 (interfaces & industrial HW)", level=1)
    doc.add_heading("A.1 General (O.13)", level=2)
    doc.add_paragraph(
        "O.13: Industrial robustness, isolation, EMC, environmental compatibility — request CE/EMC reports."
    )

    doc.add_heading("A.2 Interface architecture (O.13.1)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A2)

    doc.add_heading("A.3 External Ethernet & fibre (O.13.1.1.1)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A3)

    doc.add_heading("A.4 Local commissioning (O.13.1.1.2)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A4)

    doc.add_heading("A.5 DSO / remote logical services (O.13.1.2, O.13.1.3)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A5)

    doc.add_heading("A.6 Plant Modbus / RS485 (O.13.1.4, O.13.1.5)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A6)

    doc.add_heading("A.7 Measurement inputs (O.13.2, O.13.2.1)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A7)

    doc.add_heading("A.8 Power & time (O.13.3, O.13.5)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A8)

    doc.add_heading("A.9 Data logger — comms events (O.14)", level=2)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_A9)

    doc.add_heading("Part B — Digital I/O (PF2 actuation / GD32)", level=1)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_B)

    doc.add_heading("Part C — Annex T (cyber / segregation — not fibre PHY)", level=1)
    add_table(doc, ["Clause", "Requirement", "Question to TesPro", "Evidence"], ROWS_C)

    doc.add_heading("Response log", level=1)
    add_table(
        doc,
        ["Clause", "TesPro answer", "Evidence file", "Status PASS/PART/GAP", "Date"],
        [("", "", "", "", "")] * 5,
    )

    doc.add_heading("References (HiTEKS repo)", level=1)
    refs = [
        "knowledge-base/08-engineering/CCI_Annex_O_Extract.md",
        "knowledge-base/08-engineering/CCI_Annex_T_Extract.md",
        "knowledge-base/08-engineering/CCI_CEI_0-16_Consolidated.md",
        "Architecture/_extracted_reg_analysis/cei_0-16_consolidated_allegato_o_it.txt",
        "lab/TR400_HW_IO_CONFIG.md",
        "lab/PHASE1_LAB_SESSION_FREEZE_2026-09-22.md",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    doc.add_paragraph()
    p = doc.add_paragraph("Prepared for: TesPro supplier HW validation · HiTEKS CCLI programme")
    p.runs[0].italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
