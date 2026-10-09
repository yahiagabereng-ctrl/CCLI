#!/usr/bin/env python3
"""Generate simple lab / conformance briefing PowerPoint for HiTEKS CCLI."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "lab" / "meetings" / "2026-10-08_HiTEKS_CCLI_Lab_Conformance_Brief.pptx"

TITLE_COLOR = RGBColor(0x1A, 0x3A, 0x5C)
ACCENT = RGBColor(0x00, 0x66, 0x99)


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str], note: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)
        if line.startswith("  "):
            p.level = 1
            p.text = line.strip()
            p.font.size = Pt(18)
    if note:
        p = body.add_paragraph()
        p.text = note
        p.font.size = Pt(14)
        p.font.italic = True
        p.font.color.rgb = ACCENT


def add_two_column_slide(
    prs: Presentation, title: str, left_title: str, left: list[str], right_title: str, right: list[str]
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    def box(left_in: float, heading: str, items: list[str]) -> None:
        tx = slide.shapes.add_textbox(Inches(left_in), Inches(1.35), Inches(4.3), Inches(5.5))
        tf = tx.text_frame
        tf.word_wrap = True
        h = tf.paragraphs[0]
        h.text = heading
        h.font.bold = True
        h.font.size = Pt(22)
        h.font.color.rgb = ACCENT
        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(17)
            p.space_before = Pt(6)

    box(0.4, left_title, left)
    box(5.0, right_title, right)


def main() -> int:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "HiTEKS CCLI — Lab & Conformance Brief",
        f"DNV / accredited lab · IEC 61850-10 & IEC TS 62351-100-3\n{today} · Platform: TesPro TG544 (OpenWrt)",
    )

    add_bullet_slide(
        prs,
        "Why this meeting",
        [
            "Align on what gets tested vs what we deliver before testing",
            "Two tracks: 61850 server (protocol) + 62351-3 transport (TLS)",
            "CEI 0-16 Annex O / T — cyber + field protocols for DSO CCI",
            "Goal: quote, test plan, and document list — not firmware deep-dive",
        ],
    )

    add_bullet_slide(
        prs,
        "Product snapshot",
        [
            "Central Plant Controller (CCI) — original design, TesPro TG544 gateway",
            "DSO interface Eth_A: IEC 61850 MMS server (Annex T)",
            "Secure MMS port 3782 (TLS per IEC 62351-3)",
            "Plain MMS lab port 102 (internal bench only)",
            "Lab DUT: frozen ccli binary + ICD/CID + deployment yaml",
        ],
    )

    add_two_column_slide(
        prs,
        "Standards map (simple)",
        "Protocol & model",
        [
            "61850-6 — SCL / ICD",
            "61850-7-2 — ACSI / services (PICS)",
            "61850-8-1 — MMS mapping",
            "61850-10 Ed.2.1 — how lab tests 61850",
        ],
        "Security & grid",
        [
            "62351-3 — TLS profile (requirements)",
            "62351-100-3 — transport test cases",
            "62351-4 — MMS end-to-end (separate track)",
            "62351-9 — PKI / certificates",
            "CEI 0-16 Annex O / T — product rules",
        ],
    )

    add_bullet_slide(
        prs,
        "Documents every lab needs (both tracks)",
        [
            "PICS — what the device claims to support",
            "PIXIT — timeouts, ports, limits (e.g. IntgPd 4000 ms, port 3782)",
            "MICS — data model as ICD (61850-6); lab may build SCD",
            "TICS — known issues (often empty)",
            "Manuals + frozen hardware/firmware version",
            "61850: PICS from 7-2 Annex A · Transport: PID = PICS + PIXIT for 62351-3",
        ],
    )

    add_bullet_slide(
        prs,
        "Track A — IEC 61850-10 (server)",
        [
            "Tests MMS server behaviour — not plant logic",
            "Bench: client simulator to DUT + optional GOOSE / time master",
            "Organized in tables (sAss, sSrv, sRp, sCtl, sTm, …)",
            "Section 6.3 cyber is NOT in 61850-10 → see 62351-100-4 / 100-6",
            "Edition 2.1 (2025-07) — use for new UCA / DNV procedures",
        ],
    )

    add_bullet_slide(
        prs,
        "61850-10 — HiTEKS priority tests",
        [
            "P0 — Table 1: PICS/PIXIT/MICS match DUT version",
            "P0 — Table 2: ICD/SCL valid; Communication; reports/control",
            "P0 — sAss: association / release",
            "P0 — sSrv: browse, Get/Set (e.g. Wlim)",
            "P0 — sRp Table 14: URCB integrity ~4 s (lab TotW ~3961 ms)",
            "P0 — sCtl Table 26: control model for Wlim / WSd",
            "P1 — sTm: time sync (Annex T UTC ±100 ms)",
            "P1 — sGop: only if GOOSE claimed in PICS",
        ],
    )

    add_bullet_slide(
        prs,
        "Track B — IEC TS 62351-100-3 (TLS)",
        [
            "Tests 62351-3 integration over TCP — not raw RFC TLS alone",
            "Must run with real app: 61850 MMS on port 3782",
            "DUT as TLS server (mandatory); client tests if claimed",
            "Separate accredited report (e.g. IMQ: 62351-3:2020 + 100-3 cases)",
            "Complements 62351-4 MMS security — does not replace it",
        ],
    )

    add_bullet_slide(
        prs,
        "62351-100-3 — test layout",
        [
            "Clause 5 / Table 1 — configuration (ciphers, TLS 1.2, ports, CRL)",
            "Clause 6 / Table 2 — normal: handshake, app data, resumption, reneg",
            "Clause 6 / Table 3 — resiliency: bad/expired/revoked cert, weak cipher",
            "Clause 7 / Tables 4–5 — pass/fail matrix + formal test report",
            "Lab logs: handshake, reneg, resumption, cert checks, security events",
        ],
    )

    add_bullet_slide(
        prs,
        "PKI & TLS (lab setup)",
        [
            "Cert A — TLS channel authentication (mutual TLS on 3782)",
            "Cert B — MMS ACSE / 62351-4 end-to-end (separate from TLS cert)",
            "Lab PKI: EJBCA + Annex-aligned CA (PEMs in repo / TSP pack)",
            "OpenSSL verification gates PASS on dev PC",
            "PIXIT: secure port 3782, trust anchors, CRL/OCSP URLs",
        ],
    )

    add_two_column_slide(
        prs,
        "Status — HAVE vs TODO",
        "Ready / in progress",
        [
            "ICD/CID + lab ICD path",
            "61850-10 & 62351-100-3 extracts in RAG",
            "EJBCA PEM pack + ceremony SOP",
            "Internal P3: browse, URCB 4 s, Wlim, chrony",
            "Test Suite Pro + deploy scripts",
        ],
        "Gaps before formal lab",
        [
            "Formal 61850 PICS/PIXIT pack",
            "Formal 62351-3 PID (PICS+PIXIT)",
            "Firmware: ISO session fix for TSP Connect :3782",
            "Accredited 62351-100-4 plan (MMS security)",
            "Licensed PDF for exact table row IDs (OCR scans)",
        ],
    )

    add_bullet_slide(
        prs,
        "Current technical blocker",
        [
            "TSP Connect on :3782 can fail before ACSE if ISO Session LI/PGI wrong",
            "Symptom: presentation/session error despite correct TLS PEMs",
            "Fix documented in lab triage — needs ccli redeploy on TG544",
            "PKI can be correct while MMS session still fails — separate issues",
        ],
    )

    add_bullet_slide(
        prs,
        "Our facts — say these first",
        [
            "Server-only DUT; Ed. 2 model (parts 6, 7-x, 8-1); no SV, no client",
            "GOOSE publish exists — claim to be decided with lab",
            "MMS on 102 (bench) and 3782 (TLS 62351-3, mutual auth)",
            "URCB IntgPd 4000 ms (TotW); Wlim / WSd controls; static datasets",
            "PKI: EJBCA lab CA; Cert A (TLS) + Cert B (MMS E2E) per Annex T",
            "Firmware identified by SHA-256 of ccli + OpenWrt build ID",
        ],
    )

    add_bullet_slide(
        prs,
        "Questions 1/5 — scope & scheme",
        [
            "61850-10 edition used: Ed. 2.1 (2025) or Ed. 2? UCA procedure version?",
            "Mandatory UCA blocks for server-only; can BRCB/Log/File/SG/SV be excluded?",
            "GOOSE publisher: required now, or claim later?",
            "Deliverable: UCA IUG certificate or report — what the Italian DSO accepts",
            "Do you run 62351-100-3 and 100-4 too? One visit, same firmware?",
            "CEI 0-16 Annex O/T functional tests — in-house or referral?",
        ],
    )

    add_bullet_slide(
        prs,
        "Questions 2/5 — documents",
        [
            "PICS template: 7-2 Annex A, UCA, or lab form?",
            "PIXIT mandatory items: timeouts, max clients, IntgPd range, ctlModel",
            "MICS = ICD alone? TICS against which TISSUES date?",
            "SCD generated by lab from our ICD, or we deliver SCD + SSD?",
            "62351-3 PID template — are §8 PICS tables accepted as is?",
            "Manual: final or lab configuration guide? Units needed (1 + spare)?",
        ],
    )

    add_bullet_slide(
        prs,
        "Questions 3/5 — bench & 61850 model",
        [
            "Client simulator used — can we get it for pre-testing?",
            "Eth_A only, or all Ethernet domains? Time master: SNTP / PTP, class?",
            "Physical I/O stimulus required, or software-driven changes?",
            "Production config mandatory, or lab mode accepted?",
            "SCL schema revision; private namespaces tolerated?",
            "Table 14: all optional fields or PIXIT-enabled only? BRCB required?",
            "Table 26: which ctlModel for Wlim / WSd? Segmentation dataset size?",
        ],
    )

    add_bullet_slide(
        prs,
        "Questions 4/5 — security (62351-3 / 100-3 / 4)",
        [
            "100-3 Tables 1–5 verbatim or national scheme (e.g. IMQ IND-F-006)?",
            "62351-3 edition: 2014+AMD2 (TLS 1.2) or 2023 Ed. 2 (TLS 1.3)?",
            "Cipher list mandatory / forbidden; TLS 1.3 required?",
            "Reneg / resumption: simulated 24 h via PIXIT? CRL + OCSP + nonce?",
            "Lab PKI or our EJBCA CA? Key sizes, cert size ≤ 8192",
            "Client-station tests waived for server-only? Syslog enough for events?",
            "62351-4: compatibility or native E2E? Is 100-4 published? EST tested?",
        ],
    )

    add_bullet_slide(
        prs,
        "Questions 5/5 — process, cost, other certs",
        [
            "Lead time, days on site, price per block, re-test fee",
            "Pre-conformance run or remote pre-check offered?",
            "Fail one block — fix and re-run in same visit?",
            "Which firmware changes force re-test? Certificate validity / renewal",
            "Report language for Italian DSO; NDA before shipping DUT",
            "Referrals: 61557-12 (ISO 17025), 62443-4-1/4-2 (ISASecure), FIPS/TPM",
        ],
        note="Full table with 'why / our position / lab answer': lab/meetings/2026-10-08_Lab_Conformance_Questionnaire.md",
    )

    add_bullet_slide(
        prs,
        "What the lab will ask us — ready answers",
        [
            "Role: server only · Model Ed. 2 · Procedure Ed. 2.1",
            "ACSI: association, directory, Get/Set values, datasets, URCB, control, time",
            "Limits: max clients (PIXIT), IntgPd 1–60 s, static datasets",
            "Control: direct normal security (confirm per DO)",
            "TLS 1.2, ECDHE AES-GCM SHA256/384, RSA 2048, ≤ 8192 octets",
            "Logging: syslog RFC 5424, event store ≥ 2048",
        ],
    )

    add_bullet_slide(
        prs,
        "Decisions to record today",
        [
            "61850 scope blocks: minimum vs extended",
            "GOOSE claim: now vs later",
            "62351-3 only vs 62351-3 + 62351-4",
            "One lab vs two labs (protocol / security / environmental)",
            "PICS / PIXIT templates to use",
            "Target test date and budget approval",
        ],
    )

    add_bullet_slide(
        prs,
        "Suggested next steps",
        [
            "1. Lab returns tailored test plan + quote from draft PICS/PIXIT",
            "2. HiTEKS completes PICS (7-2) + 62351-3 PID",
            "3. Deploy session fix + EJBCA PEMs on TG544; capture TSP evidence",
            "4. Schedule 61850 server session then 62351-100-3 transport session",
            "5. Archive test report → CEI Annex O conformity file",
        ],
    )

    add_bullet_slide(
        prs,
        "Corpus references (internal)",
        [
            "knowledge-base/08-engineering/CCI_61850-10_Extract.md",
            "knowledge-base/08-engineering/CCI_62351-100-3_Extract.md",
            "knowledge-base/08-engineering/CCI_62351-3_Extract.md",
            "knowledge-base/08-engineering/CCI_Certificate_Samples_Extract.md",
            "lab/meetings/2026-10-08_Lab_Conformance_Questionnaire.md",
            "lab/EJBCA_LAB_SETUP.md · lab/evidence/",
            "RAG: ccli-61850-10-extract · ccli-62351-100-3-extract · ccli-lab-conformance-questionnaire",
        ],
        note="Regenerate: python scripts/generate-lab-conformance-deck.py",
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
