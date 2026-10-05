#!/usr/bin/env python3
"""Row-by-row compare LN,DO.xlsx mandatory DO audit vs lab_tg544_eth_a CID/cfg."""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
CID = ROOT / "apps/ccli/config/icd/lab_tg544_eth_a.cid"
CFG = ROOT / "apps/ccli/config/icd/lab_tg544_eth_a.cfg"
DEFAULT_XLSX = ROOT / "New folder/LN,DO.xlsx"
OUT_MD = ROOT / "lab/evidence/phase5/P5_LN_DO_XLSX_COMPARE.md"

# Annex T / audit shorthand → CID DO name (CEI TR 57-126 naming)
DO_ALIASES: dict[str, str] = {
    "WMaxSpt": "WMaxSptPct",
    "WSpt": "WSptPct",
    "VArTgtSpt": "VArTgtSptPct",
    "PFSpt": "PFSptPct",
    "VArSpt": "VArSptPct",
    "PFSptPct": "PFSptPct",  # identity
}


def parse_cid_lns(text: str) -> dict[str, dict]:
    lntypes: dict[str, list[str]] = {}
    for m in re.finditer(r'<LNodeType id="([^"]+)"[^>]*>(.*?)</LNodeType>', text, re.S):
        lid, body = m.groups()
        lntypes[lid] = re.findall(r'<DO name="([^"]+)"', body)

    out: dict[str, dict] = {}
    # LLN0 uses <LN0 ...> not <LN ...>
    for m in re.finditer(
        r'<LN0\s+lnClass="([^"]+)"\s+lnType="([^"]+)"\s+inst="([^"]*)"',
        text,
    ):
        cls, lntype, inst = m.groups()
        out["LLN0"] = {
            "lnClass": cls,
            "lnType": lntype,
            "prefix": "",
            "inst": inst,
            "dos": set(lntypes.get(lntype, [])),
        }
    for m in re.finditer(
        r'<LN\s+lnClass="([^"]+)"\s+lnType="([^"]+)"\s+inst="([^"]*)"\s+prefix="([^"]*)"',
        text,
    ):
        cls, lntype, inst, pref = m.groups()
        name = f"{pref}{cls}{inst}" if pref else f"{cls}{inst}"
        out[name] = {
            "lnClass": cls,
            "lnType": lntype,
            "prefix": pref,
            "inst": inst,
            "dos": set(lntypes.get(lntype, [])),
        }
    return out


def parse_cfg_lns(text: str) -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    current: str | None = None
    for line in text.splitlines():
        m = re.match(r"^LN\(([^)]+)\)\{", line)
        if m:
            current = m.group(1)
            out[current] = set()
            continue
        if current and line.startswith("DO("):
            dm = re.match(r"^DO\(([^)\s]+)", line)
            if dm:
                out[current].add(dm.group(1))
    return out


def resolve_do(do: str) -> list[str]:
    """Return candidate DO names to match in CID."""
    cands = [do]
    if do in DO_ALIASES:
        cands.append(DO_ALIASES[do])
    # suffix variants used in some audits
    if do.endswith("Spt") and do + "Pct" not in cands:
        cands.append(do + "Pct")
    return cands


def audit_row(
    ln_class: str,
    do: str,
    cid_by_class: dict[str, list[tuple[str, set[str]]]],
    cfg_dos: dict[str, set[str]],
) -> tuple[str, str, list[str]]:
    """Return (verdict, detail, missing_instances)."""
    instances = cid_by_class.get(ln_class, [])
    if not instances:
        return "NO_LN_CLASS", "lnClass not in CID", []

    cands = resolve_do(do)
    missing: list[str] = []
    present_cid = 0
    present_cfg = 0
    matched_names: set[str] = set()

    for name, dos in instances:
        hit = next((c for c in cands if c in dos), None)
        if hit:
            present_cid += 1
            matched_names.add(hit)
            if name in cfg_dos and hit in cfg_dos[name]:
                present_cfg += 1
            else:
                missing.append(f"{name}(cfg missing {hit})")
        else:
            missing.append(name)

    n = len(instances)
    if present_cid == n:
        if present_cfg == n:
            alias = ""
            if do not in matched_names and matched_names:
                alias = f" (audit `{do}` → CID `{next(iter(matched_names))}`)"
            return "MATCH", f"CID+CFG all {n} inst.{alias}", []
        return "PARTIAL_CFG", f"CID all {n}; CFG {present_cfg}/{n}", missing

    if present_cid == 0:
        # show closest DO names for first instance
        sample = sorted(instances[0][1])
        near = [d for d in sample if do.lower() in d.lower() or d.lower().startswith(do.lower()[:4])]
        hint = f"; CID has: {', '.join(near[:5])}" if near else f"; CID DOs: {', '.join(sample[:8])}"
        return "MISSING", f"0/{n} instances{hint}", missing

    return "PARTIAL", f"CID {present_cid}/{n}; CFG {present_cfg}/{n}", missing


def load_xlsx_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Mandatory DO audit"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        rows.append(
            {
                "row": i,
                "lnClass": str(row[0]).strip(),
                "section": row[1],
                "do": str(row[2]).strip(),
                "cdc": row[3],
                "std_presence": row[4],
                "std_source": row[5],
                "annex_prescond": row[6],
                "result_cci": row[7],
                "xlsx_cid": row[8],
                "xlsx_missing": row[9],
                "notes": row[10],
            }
        )
    wb.close()
    return rows


def main() -> None:
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        print(f"XLSX not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    cid_lns = parse_cid_lns(CID.read_text(encoding="utf-8", errors="replace"))
    cfg_dos = parse_cfg_lns(CFG.read_text(encoding="utf-8", errors="replace"))

    cid_by_class: dict[str, list[tuple[str, set[str]]]] = defaultdict(list)
    for name, info in sorted(cid_lns.items()):
        cid_by_class[info["lnClass"]].append((name, info["dos"]))

    xlsx_rows = load_xlsx_rows(xlsx_path)

    lines: list[str] = [
        "# P5 — LN,DO.xlsx row-by-row compare vs `lab_tg544_eth_a`",
        "",
        f"**Source:** `{xlsx_path}`",
        f"**CID:** `{CID.relative_to(ROOT)}` ({len(cid_lns)} LNs)",
        f"**CFG:** `{CFG.relative_to(ROOT)}` ({len(cfg_dos)} LNs)",
        "",
        "**Verdict legend:** `MATCH` = DO on all CID instances + cfg · `PARTIAL`/`MISSING` = gaps · `NO_LN_CLASS` = lnClass absent",
        "",
        "| Row | LN class | DO | Result (CCI) | XLSX CID col | **Our verdict** | Detail | Missing instances |",
        "|-----|----------|----|--------------|--------------|-----------------|--------|-------------------|",
    ]

    counts: dict[str, int] = defaultdict(int)
    mismatches_vs_xlsx: list[str] = []

    for r in xlsx_rows:
        verdict, detail, missing = audit_row(r["lnClass"], r["do"], cid_by_class, cfg_dos)
        counts[verdict] += 1
        miss_str = ", ".join(missing[:3])
        if len(missing) > 3:
            miss_str += f" (+{len(missing) - 3} more)"
        xlsx_cid = r["xlsx_cid"] or "—"
        lines.append(
            f"| {r['row']} | `{r['lnClass']}` | `{r['do']}` | {r['result_cci']} | {xlsx_cid} | **{verdict}** | {detail} | {miss_str or '—'} |"
        )

        # flag when our audit disagrees with pre-filled xlsx column
        xcid = str(xlsx_cid).upper()
        if verdict == "MATCH" and "NO" in xcid:
            mismatches_vs_xlsx.append(f"R{r['row']} {r['lnClass']}.{r['do']}: xlsx says NO, we say MATCH ({detail})")
        elif verdict in ("MISSING", "NO_LN_CLASS") and "YES" in xcid:
            mismatches_vs_xlsx.append(f"R{r['row']} {r['lnClass']}.{r['do']}: xlsx says YES, we say {verdict}")

    lines.extend(
        [
            "",
            "## Summary",
            "",
            f"| Verdict | Count |",
            f"|---------|------:|",
        ]
    )
    for k in sorted(counts.keys()):
        lines.append(f"| {k} | {counts[k]} |")

    lines.extend(["", f"**Total rows:** {len(xlsx_rows)}", ""])

    intentional = [
        ("DGEN", "Health", "DisFRDGEN1", "Annex T: Health only on SSGG instances (SSGGDGEN1/2 have it)"),
        ("DGEN", "GnGrId", "DisFRDGEN1", "Annex T CEI extension: GnGrId only on SSGG instances"),
    ]
    partial_rows = [(r["lnClass"], r["do"]) for r in xlsx_rows]
    if any((cls, do) in [(i[0], i[1]) for i in intentional] for cls, do in partial_rows):
        lines.extend(
            [
                "## Intentional PARTIAL (regulatory — do not add to DisFRDGEN1)",
                "",
                "| LN | DO | Missing on | Reason |",
                "|----|----|-----------:|--------|",
            ]
        )
        for cls, do, miss, reason in intentional:
            lines.append(f"| `{cls}` | `{do}` | `{miss}` | {reason} |")
        lines.append("")

    if mismatches_vs_xlsx:
        lines.extend(["## Disagreements vs xlsx «In your CID?» column", ""])
        for m in mismatches_vs_xlsx:
            lines.append(f"- {m}")
        lines.append("")

    lines.extend(
        [
            "## CID instances by lnClass",
            "",
            "| lnClass | Instances |",
            "|---------|-----------|",
        ]
    )
    for cls, insts in sorted(cid_by_class.items()):
        names = ", ".join(f"`{n}`" for n, _ in insts)
        lines.append(f"| `{cls}` | {names} |")

    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT_MD} ({len(xlsx_rows)} rows)")
    print("Summary:", dict(counts))
    if mismatches_vs_xlsx:
        print(f"Disagreements vs xlsx column: {len(mismatches_vs_xlsx)}")


if __name__ == "__main__":
    main()
