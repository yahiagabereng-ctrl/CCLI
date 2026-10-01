#!/usr/bin/env python3
"""Generate lab/CCI_Figura2_Parameters.docx and copy to Desktop (if path exists)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_REPO = REPO_ROOT / "lab" / "CCI_Figura2_Parameters.docx"
OUT_DESKTOP = Path(r"c:\Users\yahia\OneDrive\Desktop\CCI_Figura2_Parameters.docx")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
    doc.add_paragraph()


def build() -> Document:
    doc = Document()
    title = doc.add_heading("Figura 2 (PF2) — Parameters, equations & CCLI mapping", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("Document ID: ").bold = True
    p.add_run("CCLI-LAB-FIG2-PAR-001  |  Rev 1.0  |  2026-09-22\n")
    p.add_run("Sources: ").bold = True
    p.add_run(
        "CEI TR 57-126 §6.2 Table 1 · CEI 0-16 Allegato O/T · "
        "CCLI apps/ccli/core/dso/ · lab/CCI_Figura2_Parameters.md"
    )

    doc.add_heading("Table numbering (Italian consolidated vs English extract)", level=1)
    add_table(
        doc,
        ["Function", "Italian Allegato T", "English CCI_Annex_T_Extract §5"],
        [
            ["Wlim (O.9.2.2)", "Tab. 84", "Table 85"],
            ["WSd (O.9.2.3)", "Tab. 85", "Table 86"],
            ["WSa (MSD)", "Tab. 86", "Table 87"],
        ],
    )

    doc.add_heading("1. Figura 2 use case (TR Table 1 — ON/OFF)", level=1)
    add_table(
        doc,
        ["Function", "TR state", "Clause", "Figura 2"],
        [
            ["Wlim", "Not operative / Inactive", "O.9.2.2", "Off"],
            ["WSd", "Operative / Active", "O.9.2.3", "On — green s.p. W"],
            ["WSptPct", "+20 % Smax", "O.9.2.3", "Export set-point"],
            ["VArSd, PFSP, VArV, PFW", "Not operative", "O.9.1.x", "Off"],
        ],
    )

    doc.add_heading("2. TR Table 1 — plant envelope (O.8.2)", level=1)
    add_table(
        doc,
        ["Parameter", "TR value", "Unit", "CCLI yaml", "Code"],
        [
            ["Max export P (Pfed)", "200", "kW", "plant.p_imm_kw", "plant_envelope.hpp"],
            ["Max import P (Pass)", "200", "kW", "plant.p_ass_kw", "plant_envelope.hpp"],
            ["Max inductive Q", "50", "kVAr", "plant.q_ind_kvar", "plant_envelope.hpp"],
            ["Max capacitive Q", "50", "kVAr", "plant.q_cap_kvar", "plant_envelope.hpp"],
            ["Smax nameplate", "210", "kVA", "plant.smax_kva", "smax_kva_effective()"],
            ["IED name", "CCI016_01", "—", "—", "config/icd/cei-tr-57-126-example.cid"],
            ["Eth_A IP (example)", "192.168.8.167", "—", "interfaces.eth_a (Phase 2)", "lab_tr400.yaml"],
        ],
    )

    doc.add_heading("3. TR Table 1 — active power (Figura 2)", level=1)
    add_table(
        doc,
        ["Parameter", "TR value", "CCLI yaml / preset"],
        [
            ["Wlim", "Inactive", "dso.wlim_active: false"],
            ["WSd", "Active", "dso.wsd_active: true (tr57126_table1_dso)"],
            ["WSptPct", "+20 %", "wspt_pct: 20"],
        ],
    )

    doc.add_heading("4. Equations & numeric results", level=1)
    doc.add_paragraph(
        "(1) Smax_calc = sqrt(max(Pimm²,Pass²)+max(Qind²,Qcap²)) → 206.15 kVA for TR corners."
    )
    doc.add_paragraph("(1b) Authoritative Smax = 210 kVA (Operating Rule / DPCC3.PdC_VA.VARtg).")
    doc.add_paragraph("(2) P_kW = (pct/100)×Smax → WSd 20% × 210 = 42 kW.")
    doc.add_paragraph("(5)(6) P_effective = min(active caps) → 42 kW (Wlim off, WSd only).")
    doc.add_paragraph(
        "(13)–(17) PF2: apply_dso_to_pf2_config → threshold 42 kW, release 37 kW "
        "(release_delta_kw: 5 in lab_tr400_dso_tr57126.yaml)."
    )
    doc.add_paragraph(
        "Gates: dso.enabled AND pf2.use_dso_mock. Measured P: Modbus reg 40001 (kW), not % mock."
    )

    doc.add_heading("5. O.7 timing", level=1)
    add_table(
        doc,
        ["Parameter", "CEI", "Phase 1 CCLI"],
        [
            ["TsP active P", "≤ 60 s, ±5 %", "settled_within_band_kw() unit only"],
            ["TsQ reactive", "≤ 10 s", "Not impl."],
            ["ΔT slow ring", "10–600 s, default 60 s", "Not impl."],
            ["Set-point spacing", "≥ 3 s (O.7.3.3)", "Not impl."],
            ["TotW period", "4 s (O.8.3)", "Modbus poll 1 s"],
        ],
    )

    doc.add_heading("6. Annex T — Wlim (English Table 85 / IT Tab. 84)", level=1)
    add_table(
        doc,
        ["Parameter", "Range", "Default", "Figura 2 TR"],
        [
            ["Operating status", "1=Op / 5=Not-Op", "5", "5 (off)"],
            ["Gen limit", "0..100 % Smax", "0", "inactive"],
            ["Activation", "5=Inactive / 1=Active", "5", "Inactive"],
        ],
    )
    doc.add_paragraph("O.9.2.1 trigger: V near 110 % Un — qualitative, no numeric equation.")

    doc.add_heading("7. Annex T — WSd (English Table 86 / IT Tab. 85)", level=1)
    add_table(
        doc,
        ["Parameter", "Range", "Default", "Figura 2 TR"],
        [
            ["Operating status", "1=Op / 5=Not-Op", "5", "1 (on)"],
            ["WSptPct", "±100 % Smax", "100/0", "+20 %"],
            ["Activation", "5=Inactive / 1=Active", "5", "Active"],
            ["Function status", "0=N/A / 2=Automatic", "2", "Automatic"],
        ],
    )

    doc.add_heading("8. Annex T — reactive (Tables 88–92) — Figura 2 all OFF", level=1)
    add_table(
        doc,
        ["Function", "Clause", "Key trigger / default", "TR Table 1"],
        [
            ["VArSd / s.p. Q", "O.9.1.4", "TsQ ≤ 10 s", "Inactive"],
            ["PFSP / s.p. cosφ", "O.9.1", "cosφ defaults ±0.95", "Inactive"],
            ["VArV / Q(V)", "O.9.1.3", "δQ = 5% Qmax", "Inactive"],
            ["PFW / cosφ(P)", "O.9.1.2", "δcosφ = 0.02", "Inactive"],
        ],
    )

    doc.add_heading("9. CCLI code reference", level=1)
    add_table(
        doc,
        ["Topic", "Path"],
        [
            ["Regulation map", "apps/ccli/core/dso/regulation_map.hpp"],
            ["Plant Eqs (1)(2)", "apps/ccli/core/dso/plant_envelope.hpp"],
            ["DSO combine O.11", "apps/ccli/core/dso/dso_active_power.cpp"],
            ["Yaml → PF2 kW", "apps/ccli/core/dso/dso_phase1.cpp"],
            ["Config parse", "apps/ccli/core/config/ccli_config.cpp"],
            ["Main loop", "apps/ccli/services/ccli_main.cpp"],
            ["PF2 FSM", "apps/ccli/core/pf2/pf2_fsm.cpp"],
            ["Modbus P", "apps/ccli/adapters/modbus_master/modbus_adapter.cpp"],
            ["TR yaml profile", "apps/ccli/config/lab_tr400_dso_tr57126.yaml"],
            ["Unit tests", "apps/ccli/tests/unit/dso_phase1_test.cpp"],
            ["Master map", "lab/PF2_REGULATION_PHASE1.md"],
        ],
    )

    doc.add_heading("10. Phase 1 limitations", level=1)
    for item in [
        "WSd → export ceiling for P>threshold FSM; not full O.9.2.3 closed-loop.",
        "Negative WSptPct (import) not mapped.",
        "Wlim/WSd from yaml/preset, not MMS Eth_A.",
        "dso.enabled without pf2.use_dso_mock does not rewrite PF2 kW.",
        "Reactive functions not in core/dso/.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading(
        "Appendix — Full Annex T matrix (Figura 2 diagram scope, English extract §5)",
        level=1,
    )
    doc.add_paragraph(
        "Detailed per-function parameters, ranges, defaults, and O.7 timestamps "
        "(Tables 85–92). Figura 2 TR instance values in §2–§3 above."
    )

    appendix_rows = [
        [
            "lim W110 / lim W (O.9.2.1/2)",
            "TsP ≤ 60 s",
            "Op status 1/5",
            "Gen limit 0..100% Smax",
            "Default 5 / 0",
            "Table 85 EN",
        ],
        [
            "Active power modulation WSd (O.9.2.3)",
            "TsP ≤ 60 s; update ≥ 3 s",
            "Op status 1/5",
            "Setpoint ±100% Smax",
            "Default 5; TR +20%",
            "Table 86 EN",
        ],
        [
            "s.p. W WSa (O.10.3.1 MSD)",
            "TsP ≤ 60 s; update ≥ 3 s",
            "Op status 1/5",
            "Setpoint ±100% Smax",
            "Default 5",
            "Table 87 EN",
        ],
        [
            "VArSd (O.9.1.4)",
            "TsQ ≤ 10 s; update ≥ 3 s",
            "Op status 1/5",
            "Q ±100% Smax",
            "Default 0",
            "Table 88 EN",
        ],
        [
            "s.p. Q (O.10.3.2)",
            "TsQ ≤ 10 s",
            "Same as VArSd",
            "—",
            "Inactive TR",
            "Table 89 EN",
        ],
        [
            "s.p. cosφ PFSP (O.9.1)",
            "TsQ ≤ 10 s",
            "cosφ −1..+1",
            "Default −0.95 / +0.95",
            "Inactive TR",
            "Table 90 EN",
        ],
        [
            "Q(V) VArV (O.9.1.3)",
            "ΔT 10–600 s; TsQ ≤ 10 s",
            "K, lock-in/out, V bands",
            "δQ = 5% Qmax",
            "Inactive TR",
            "Table 91 EN",
        ],
        [
            "cosφ(P) PFW (O.9.1.2)",
            "ΔT 10–600 s; TsQ ≤ 10 s",
            "Points A/B/C P and cosφ",
            "δcosφ = 0.02",
            "Inactive TR",
            "Table 92 EN",
        ],
    ]
    add_table(
        doc,
        ["Function", "Timestamp O.7", "Key params", "Notes", "Figura 2", "Annex T"],
        appendix_rows,
    )

    return doc


def main() -> None:
    doc = build()
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_REPO)
    print(f"Wrote {OUT_REPO}")
    if OUT_DESKTOP.parent.is_dir():
        shutil.copy2(OUT_REPO, OUT_DESKTOP)
        print(f"Copied to {OUT_DESKTOP}")
    else:
        print(f"Desktop path not found: {OUT_DESKTOP.parent}")


if __name__ == "__main__":
    main()
